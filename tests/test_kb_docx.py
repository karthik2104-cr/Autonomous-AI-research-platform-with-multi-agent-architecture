"""Knowledge Base DOCX ingest → embed → query test."""
from __future__ import annotations

from pathlib import Path

import tests  # noqa: F401 — repo root on sys.path
from tests import ROOT

from core.ingest import ingest_file, parse_file
from core.vector_store import document_count, list_sources, query_documents, reset_collection

SAMPLE_DOCX = ROOT / "testdata" / "helios_brief.docx"


def _ensure_sample_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph(
        "Project Helios is an internal research codename for edge caching experiments "
        "on lab hardware."
    )
    doc.add_paragraph(
        "The team stores experiment notes in this document so the knowledge base can "
        "retrieve them during research runs."
    )
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Component"
    table.rows[0].cells[1].text = "Latency target"
    table.rows[1].cells[0].text = "Edge cache"
    table.rows[1].cells[1].text = "Under 12 milliseconds"
    doc.add_paragraph("Helios deployments must stay offline — no cloud API calls.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main() -> None:
    if not SAMPLE_DOCX.exists():
        print(f"Creating sample DOCX at {SAMPLE_DOCX}")
        _ensure_sample_docx(SAMPLE_DOCX)

    print("=== PARSE DOCX ===")
    text, filename = parse_file(SAMPLE_DOCX)
    print("filename:", filename)
    print("extracted_text_length:", len(text))
    print("extracted_preview:", repr(text[:320]))
    assert text.strip(), "Expected non-empty DOCX text"
    assert "Project Helios" in text
    assert "Edge cache" in text and "12 milliseconds" in text

    reset_collection()
    print("\n=== INGEST DOCX ===")
    result = ingest_file(SAMPLE_DOCX)
    print("ingest_result:", result)
    print("chunk_count:", document_count())
    print("sources:", list_sources())
    assert result["filename"] == "helios_brief.docx"
    assert result["chunks"] >= 1

    print("\n=== QUERY: edge cache latency ===")
    hits = query_documents("What is the latency target for the edge cache?", n_results=3)
    for i, hit in enumerate(hits, start=1):
        print(f"\n--- hit {i} ---")
        print("filename:", hit.get("filename"))
        print("chunk_id:", hit.get("chunk_id"))
        print("chunk_index:", hit.get("chunk_index"))
        print("doc_id:", hit.get("doc_id"))
        print("distance:", hit.get("distance"))
        print("text:", (hit.get("text") or "")[:400])

    assert hits, "Expected at least one retrieved chunk"
    joined = " ".join(h.get("text") or "" for h in hits).lower()
    assert "12" in joined or "millisecond" in joined, f"Expected latency table text, got: {joined[:200]}"
    assert any(h.get("filename") == "helios_brief.docx" for h in hits)
    print("\nKB_DOCX_E2E_OK")


if __name__ == "__main__":
    main()
